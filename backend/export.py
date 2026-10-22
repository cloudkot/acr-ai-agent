import json
import asyncio
import time
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches, Pt as PptxPt
from schemas import Framework
from logger import logger

def _export_to_docx_sync(framework: Framework, content: dict, filename: str, constraints_check: dict = None, generation_time: float = 0.0):
    doc = Document()
    
    # --- 1. Титульный лист (по шаблону АЦР) ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("АВТОНОМНАЯ НЕКОММЕРЧЕСКАЯ ОРГАНИЗАЦИЯ\nДОПОЛНИТЕЛЬНОГО ПРОФЕССИОНАЛЬНОГО ОБРАЗОВАНИЯ\n«АКАДЕМИЯ ИНФОРМАЦИОННЫХ ТЕХНОЛОГИЙ»")
    run.bold = True
    
    doc.add_paragraph("\n\n\n\n")
    
    order_para = doc.add_paragraph()
    order_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    order_para.add_run("УТВЕРЖДЕНО\nПриказом № ______-ОБР\nот «___» ____________ 2026 г.\n\n")
    
    prog_type = "ДОПОЛНИТЕЛЬНАЯ ПРОФЕССИОНАЛЬНАЯ ПРОГРАММА\n" if framework.template_type == "ДПО" else "ДОПОЛНИТЕЛЬНАЯ ОБЩЕРАЗВИВАЮЩАЯ ПРОГРАММА\n"
    prog_title = doc.add_paragraph()
    prog_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prog_title_run = prog_title.add_run(prog_type + f"«{framework.goal[:100]}»\n\nОбъем: {framework.total_hours} академических часов")
    prog_title_run.bold = True
    
    doc.add_paragraph("\n\n\n\n")
    doc.add_paragraph("Магас\n2026").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # --- 2. Цель и нормативная база ---
    doc.add_heading('1. Общая характеристика программы', level=1)
    doc.add_heading('1.1. Цель программы', level=2)
    doc.add_paragraph(framework.goal)
    
    if framework.prof_standard:
        doc.add_paragraph(f"Нормативная база: {framework.prof_standard}")

    if constraints_check and not constraints_check.get("constraints_ok", True):
        doc.add_heading('ВНИМАНИЕ: Предупреждения по ограничениям', level=2)
        p = doc.add_paragraph()
        p.runs[0].font.color.rgb = RGBColor(255, 0, 0)
        p.add_run("Нейросеть не до конца выполнила требования методиста:\n")
        for err in constraints_check.get("errors", []):
            doc.add_paragraph(err, style='List Bullet')

    # --- 3. Карта компетенций ---
    doc.add_heading('2. Планируемые результаты обучения (Карта компетенций)', level=1)
    
    table_comp = doc.add_table(rows=1, cols=3)
    table_comp.style = 'Table Grid'
    hdr = table_comp.rows[0].cells
    hdr[0].text = "Компетенция"
    hdr[1].text = "Знания (что должен знать)"
    hdr[2].text = "Умения и навыки (что должен уметь)"
    for cell in hdr:
        cell.paragraphs[0].runs[0].bold = True

    for comp in framework.competencies:
        row = table_comp.add_row().cells
        row[0].text = comp.name
        row[1].text = "\n".join([f"- {k}" for k in comp.knowledge])
        row[2].text = "\n".join([f"- {s}" for s in comp.skills])

    # --- 4. Учебный план ---
    doc.add_heading('3. Учебный план', level=1)
    
    table_plan = doc.add_table(rows=1, cols=7)
    table_plan.style = 'Table Grid'
    hdr_plan = table_plan.rows[0].cells
    headers = ["№", "Наименование модуля", "Всего ч.", "ЛК", "ПР", "ЛБ", "СРС"]
    for i, text in enumerate(headers):
        hdr_plan[i].text = text
        hdr_plan[i].paragraphs[0].runs[0].bold = True

    for i, mod in enumerate(framework.modules, 1):
        row = table_plan.add_row().cells
        row[0].text = str(i)
        row[1].text = mod.title
        row[2].text = str(mod.total_hours)
        row[3].text = str(mod.lectures)
        row[4].text = str(mod.practice)
        row[5].text = str(mod.labs)
        row[6].text = str(mod.srs)

    row = table_plan.add_row().cells
    row[1].text = "Итоговая аттестация"
    row[2].text = str(framework.final_assessment_hours)
    for i in [0, 3, 4, 5, 6]: row[i].text = "-"

    row = table_plan.add_row().cells
    row[1].text = "ИТОГО"
    row[1].paragraphs[0].runs[0].bold = True
    row[2].text = str(framework.total_hours)
    row[2].paragraphs[0].runs[0].bold = True

    # --- 5. Рабочие программы ---
    if content.get("work_programs"):
        doc.add_heading('4. Рабочие программы модулей', level=1)
        for wp in content["work_programs"]:
            doc.add_heading(wp.get("module_title", "Модуль"), level=2)
            doc.add_paragraph(wp.get("description", ""))

    # --- 6. Тесты ---
    doc.add_heading('5. Фонд оценочных средств (Тесты)', level=1)
    for i, test in enumerate(content.get("tests", []), 1):
        doc.add_paragraph(f"Вопрос {i}: {test.get('question')}")
        for opt in test.get('options', []):
            prefix = "[+] " if opt == test.get('correct_answer') else "[ ] "
            doc.add_paragraph(f"{prefix}{opt}", style='List Bullet')
        doc.add_paragraph(f"Проверяемая компетенция: {test.get('related_competency', 'Не указана')}\n")

    # --- 7. Семинары ---
    doc.add_heading('6. Задания для семинарских занятий', level=1)
    for i, sem in enumerate(content.get("seminars", []), 1):
        doc.add_heading(f"Семинар {i}: {sem.get('module_title', '')}", level=2)
        doc.add_paragraph(f"Задание: {sem.get('task')}")
        doc.add_paragraph(f"Критерии оценки: {sem.get('criteria')}")

    # --- 8. Лекции ---
    doc.add_heading('7. Планы лекций', level=1)
    for lec in content.get("lectures", []):
        doc.add_heading(lec.get('module_title', 'Модуль'), level=2)
        for thesis in lec.get('theses', []):
            doc.add_paragraph(thesis, style='List Bullet')

    # --- 9. Библиография ---
    if content.get("bibliography"):
        doc.add_heading('8. Список литературы', level=1)
        for book in content["bibliography"]:
            doc.add_paragraph(book, style='List Number')

    # --- 10. Кадровое обеспечение (Пустые поля) ---
    doc.add_page_break()
    doc.add_heading('9. Кадровое обеспечение программы', level=1)
    doc.add_paragraph("ФИО преподавателя: _______________________________________")
    doc.add_paragraph("Ученая степень/звание: ___________________________________")
    doc.add_paragraph("Должность: _______________________________________________")
    doc.add_paragraph("\n\nРазработчик программы: ____________________ / _________________ /")
    doc.add_paragraph("Согласовано: ____________________ / _________________ /")
    
    doc.add_paragraph(f"\n\n[Сгенерировано автоматически за {generation_time:.2f} секунд]").runs[0].font.size = Pt(8)

    doc.save(filename)
    return filename

def _export_to_pptx_sync(framework: Framework, content: dict, filename: str, constraints_check: dict = None):
    prs = Presentation()
    
    # 1. Титульный слайд
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"Программа: {framework.goal[:50]}..."
    slide.placeholders[1].text = f"Объем: {framework.total_hours} ч.\nТип: {framework.template_type}"
    
    # 2. Слайд с предупреждениями (если есть)
    if constraints_check and not constraints_check.get("constraints_ok", True):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Внимание: Отклонения от ТЗ"
        tf = slide.placeholders[1].text_frame
        tf.text = "При генерации были нарушены ограничения:"
        for err in constraints_check.get("errors", []):
            p = tf.add_paragraph()
            p.text = err
            p.level = 1

    # 3. Слайды из контента (используем новые типы слайдов)
    slides_data = content.get("slides", [])
    
    # Если слайдов нет, делаем fallback из лекций
    if not slides_data and content.get("lectures"):
        for lec in content["lectures"]:
            slides_data.append({
                "slide_type": "content",
                "title": lec.get("module_title", "Лекция"),
                "bullets": lec.get("theses", []),
                "speaker_notes": "Описание лекции отсутствует."
            })

    for slide_info in slides_data:
        slide_type = slide_info.get("slide_type", "content")
        title_text = slide_info.get("title", "Слайд")
        bullets = slide_info.get("bullets", [])
        notes_text = slide_info.get("speaker_notes", "")

        # Выбор макета
        if slide_type == "title":
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title_text
            if bullets:
                slide.placeholders[1].text = bullets[0]
        elif slide_type == "two_columns" and len(bullets) >= 2:
            slide = prs.slides.add_slide(prs.slide_layouts[3]) # Two Content
            slide.shapes.title.text = title_text
            # Заполняем две колонки
            tf_left = slide.placeholders[1].text_frame
            tf_right = slide.placeholders[2].text_frame
            mid = len(bullets) // 2
            tf_left.text = bullets[0]
            for b in bullets[1:mid]: tf_left.add_paragraph(b)
            tf_right.text = bullets[mid]
            for b in bullets[mid+1:]: tf_right.add_paragraph(b)
        else:
            # Обычный слайд (Title and Content)
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title_text
            tf = slide.placeholders[1].text_frame
            if bullets:
                tf.text = bullets[0]
                for bullet in bullets[1:]:
                    p = tf.add_paragraph()
                    p.text = bullet
                    p.level = 1
                    
        # Добавление заметок для лектора
        if notes_text:
            slide.notes_slide.notes_text_frame.text = notes_text

    prs.save(filename)
    return filename

def _export_tests_to_lms_sync(tests_data: list, filename: str):
    with open(filename, 'w', encoding='utf-8') as f:
        for i, test in enumerate(tests_data, 1):
            question = test.get('question', f'Вопрос {i}')
            options = test.get('options', [])
            correct = test.get('correct_answer', '')
            f.write(f"::Вопрос {i}::{question} {{\n")
            for opt in options:
                if opt == correct: f.write(f"  ={opt}\n")
                else: f.write(f"  ~{opt}\n")
            f.write("}\n\n")
    return filename

def _export_tests_to_moodle_xml_sync(tests_data: list, filename: str):
    quiz = ET.Element("quiz")
    for i, test in enumerate(tests_data, 1):
        question = ET.SubElement(quiz, "question", type="multichoice")
        name = ET.SubElement(question, "name")
        ET.SubElement(name, "text").text = f"Вопрос {i}"
        
        q_text = ET.SubElement(question, "questiontext", format="html")
        ET.SubElement(q_text, "text").text = f"<p>{test.get('question', '')}</p>"
        
        for opt in test.get('options', []):
            fraction = "100" if opt == test.get('correct_answer') else "0"
            answer = ET.SubElement(question, "answer", fraction=fraction, format="html")
            ET.SubElement(answer, "text").text = f"<p>{opt}</p>"
            
    tree = ET.ElementTree(quiz)
    ET.indent(tree, space="  ")
    tree.write(filename, encoding="utf-8", xml_declaration=True)
    return filename

def _export_to_json_sync(framework: Framework, content: dict, filename: str, constraints_check: dict = None, generation_time: float = 0.0):
    data = {
        "framework": framework.model_dump(),
        "content": content,
        "constraints_check": constraints_check,
        "metrics": {"generation_time_seconds": generation_time}
    }
    with open(filename, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)
    return filename

async def export_to_docx(framework: Framework, content: dict, filename: str, constraints_check: dict = None, generation_time: float = 0.0):
    return await asyncio.to_thread(_export_to_docx_sync, framework, content, filename, constraints_check, generation_time)

async def export_to_pptx(framework: Framework, content: dict, filename: str, constraints_check: dict = None):
    return await asyncio.to_thread(_export_to_pptx_sync, framework, content, filename, constraints_check)

async def export_tests_to_lms(tests_data: list, filename: str):
    return await asyncio.to_thread(_export_tests_to_lms_sync, tests_data, filename)

async def export_tests_to_moodle_xml(tests_data: list, filename: str):
    return await asyncio.to_thread(_export_tests_to_moodle_xml_sync, tests_data, filename)

async def export_to_json(framework: Framework, content: dict, filename: str, constraints_check: dict = None, generation_time: float = 0.0):
    return await asyncio.to_thread(_export_to_json_sync, framework, content, filename, constraints_check, generation_time)